import hashlib
import math
import os
import re
from types import SimpleNamespace

import pytest


@pytest.fixture
def isolated_env(monkeypatch, tmp_path):
    saved = dict(os.environ)
    for key in list(os.environ):
        if key.startswith("RAG_"):
            del os.environ[key]
    monkeypatch.chdir(tmp_path)
    yield tmp_path
    os.environ.clear()
    os.environ.update(saved)


class FakeOllamaClient:
    def __init__(self, models=(), error=None, reply="An answer [1]."):
        self.models = models
        self.error = error
        self.reply = reply
        self.chats = []

    def list(self):
        if self.error:
            raise self.error
        return SimpleNamespace(models=[SimpleNamespace(model=m) for m in self.models])

    def chat(self, **kwargs):
        self.chats.append(kwargs)
        for word in self.reply.split(" "):
            yield SimpleNamespace(message=SimpleNamespace(content=word + " "))


class FakeEmbedder:
    """Bag-of-words hashing: texts sharing words get similar vectors."""

    dim = 64

    def __init__(self, max_tokens=256):
        self.max_tokens = max_tokens

    def embed(self, texts):
        return [self._vector(t) for t in texts]

    def count_tokens(self, text):
        return len(text.split())

    def _vector(self, text):
        v = [0.0] * self.dim
        for word in re.findall(r"\w+", text.lower()):
            v[int(hashlib.md5(word.encode()).hexdigest(), 16) % self.dim] += 1.0  # noqa: S324
        norm = math.sqrt(sum(x * x for x in v)) or 1.0
        return [x / norm for x in v]


@pytest.fixture
def make_docx():
    from docx import Document

    def _make(path, *paragraphs, table=None):
        path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
        if table:
            t = doc.add_table(rows=len(table), cols=len(table[0]))
            for r, row in enumerate(table):
                for c, value in enumerate(row):
                    t.cell(r, c).text = value
        doc.save(path)
        return path

    return _make
