from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

MAX_FILE_BYTES = 25 * 1024 * 1024

_SPACES = re.compile(r"[ \t ]+")
# "high heat.In a bowl" -> "high heat. In a bowl"
_MISSING_SPACE = re.compile(r"(?<=[a-z0-9)])([.!?])(?=[A-Z])")


@dataclass(frozen=True)
class Recipe:
    id: str
    title: str
    text: str
    files: tuple[str, ...]


@dataclass
class LoadReport:
    recipes: list[Recipe] = field(default_factory=list)
    skipped: list[tuple[str, str]] = field(default_factory=list)


def load_recipes(root: Path) -> LoadReport:
    """Each folder with .docx files becomes one recipe, titled by the folder name."""
    report = LoadReport()
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(d for d in dirnames if not d.startswith("."))
        folder = Path(dirpath)
        texts, used = [], []
        for name in sorted(filenames):
            if not name.lower().endswith(".docx") or name.startswith((".", "~$")):
                continue
            path = folder / name
            rel = path.relative_to(root).as_posix()
            reason = _skip_reason(path)
            if reason is None:
                try:
                    text = read_docx(path)
                except Exception as exc:
                    reason = f"unreadable ({type(exc).__name__})"
                else:
                    reason = None if text else "no text"
            if reason:
                report.skipped.append((rel, reason))
                continue
            texts.append(text)
            used.append(rel)
        if texts:
            report.recipes.append(
                Recipe(
                    id=folder.relative_to(root).as_posix(),
                    title=clean_title(folder.name),
                    text="\n".join(texts),
                    files=tuple(used),
                )
            )
    return report


def read_docx(path: Path) -> str:
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(filter(None, (clean_line(line) for line in lines)))


def clean_line(line: str) -> str:
    return _MISSING_SPACE.sub(r"\1 ", _SPACES.sub(" ", line)).strip()


def clean_title(name: str) -> str:
    return _SPACES.sub(" ", name).strip()


def _skip_reason(path: Path) -> str | None:
    if path.is_symlink():
        return "symlink"
    if path.stat().st_size > MAX_FILE_BYTES:
        return "too large"
    return None
